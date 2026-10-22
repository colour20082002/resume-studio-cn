# -*- coding: utf-8 -*-
"""从 resume.yml 生成 PDF（Edge 无头打印）与可编辑 Word。"""
import argparse
import subprocess
import sys
from pathlib import Path

import yaml
from jinja2 import Environment, FileSystemLoader

SCRIPT_DIR = Path(__file__).resolve().parent      # 模板与版式资源
ROOT = Path.cwd()                                 # 简历内容目录（resume.yml / photo.png）
OUT_DIR = ROOT                                    # 产物目录
BUILD = ROOT / "build"

DATA = {}
EDGE = Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")

DOCX_OUT = OUT_DIR / "个人简历.docx"

ENV = Environment(loader=FileSystemLoader(str(SCRIPT_DIR)), autoescape=False,
                  keep_trailing_newline=True)


def configure(root, out=None) -> None:
    """切换简历内容目录与产物目录，重新加载 resume.yml"""
    global ROOT, OUT_DIR, BUILD, DATA, DOCX_OUT
    ROOT = Path(root).resolve()
    OUT_DIR = Path(out).resolve() if out else ROOT
    BUILD = ROOT / "build"
    BUILD.mkdir(parents=True, exist_ok=True)
    DATA = yaml.safe_load((ROOT / "resume.yml").read_text(encoding="utf-8"))
    DOCX_OUT = OUT_DIR / "个人简历.docx"

# 版式：key -> (中文名, html 模板, 版式 css)
LAYOUTS = {
    "classic": ("经典", "classic.html.j2", "classic.css"),
    "sidebar": ("左栏", "sidebar.html.j2", "sidebar.css"),
    "band": ("色带", "band.html.j2", "band.css"),
    "timeline": ("时间轴", "timeline.html.j2", "timeline.css"),
    "grid": ("卡片", "grid.html.j2", "grid.css"),
    "centered": ("居中", "centered.html.j2", "centered.css"),
}

# 配色：key -> (中文名, 配色 css, 主色，供编辑器色块使用)
PALETTES = {
    "ink": ("墨黑", "ink.css", "#1B2228"),
    "gold": ("鎏金", "gold.css", "#C0A265"),
    "navy": ("黛蓝", "navy.css", "#22405F"),
    "burgundy": ("绛红", "burgundy.css", "#7B2E3C"),
    "forest": ("墨绿", "forest.css", "#2E5A48"),
    "slate": ("烟灰", "slate.css", "#6C7681"),
}

# 字体组合：key -> (中文名, 字体 css)
FONTS = {
    "songhei": ("宋黑对照", "songhei.css"),
    "quansong": ("全宋典雅", "quansong.css"),
    "yahei": ("雅黑现代", "yahei.css"),
    "dengxian": ("等线清爽", "dengxian.css"),
    "editorial": ("黑标宋文", "editorial.css"),
    "lightsans": ("细黑雅致", "lightsans.css"),
}

# 默认 6 个组合（版式, 配色, 字体），也是母版批量生成的清单
MATCHES = [
    ("classic", "ink", "songhei"),
    ("sidebar", "navy", "songhei"),
    ("band", "forest", "yahei"),
    ("timeline", "burgundy", "songhei"),
    ("grid", "gold", "dengxian"),
    ("centered", "slate", "quansong"),
]


def pdf_path(layout: str, palette: str, font: str) -> Path:
    return OUT_DIR / (f"个人简历-{LAYOUTS[layout][0]}-{PALETTES[palette][0]}"
                      f"-{FONTS[font][0]}.pdf")


def render_html(layout: str, palette: str, font: str) -> Path:
    photo = ROOT / "photo.png"
    if photo.exists():
        (BUILD / "photo.png").write_bytes(photo.read_bytes())
    d = {**DATA, "show_photo": bool(DATA.get("show_photo", True)) and photo.exists()}
    html = ENV.get_template("template.html.j2").render(
        d=d, layout=layout, palette=palette, font=font,
        base_css=(SCRIPT_DIR / "base.css").read_text(encoding="utf-8"),
        layout_css=(SCRIPT_DIR / "layouts" / LAYOUTS[layout][2]).read_text(encoding="utf-8"),
        palette_css=(SCRIPT_DIR / "palettes" / PALETTES[palette][1]).read_text(encoding="utf-8"),
        font_css=(SCRIPT_DIR / "fonts" / FONTS[font][1]).read_text(encoding="utf-8"),
    )
    out = BUILD / f"resume-{layout}-{palette}-{font}.html"
    out.write_text(html, encoding="utf-8")
    return out


def edge(*args: str) -> None:
    tmp = BUILD / "profile"
    cmd = [str(EDGE), "--headless=new", "--disable-gpu", "--hide-scrollbars",
           f"--user-data-dir={tmp}", *args]
    r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="ignore")
    if r.returncode != 0:
        sys.stderr.write(r.stdout + r.stderr)
        raise SystemExit(f"edge failed: {r.returncode}")


def make_pdf(html: Path, out: Path) -> None:
    if out.exists():
        out.unlink()
    edge("--no-pdf-header-footer", f"--print-to-pdf={out}", html.as_uri())


def make_png(html: Path, layout: str, palette: str, font: str) -> None:
    edge("--window-size=820,3400",
         f"--screenshot={BUILD / f'preview-{layout}-{palette}-{font}.png'}", html.as_uri())


# ------------------------------- Word -------------------------------
def make_docx() -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    FONT = "微软雅黑"
    INK = RGBColor(0x11, 0x11, 0x11)
    BODY = RGBColor(0x3D, 0x3D, 0x3D)
    GRAY = RGBColor(0x77, 0x77, 0x77)
    CONTENT_W = 18.0

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.3), Cm(1.2)
    sec.left_margin = sec.right_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    def para(space_after=0, space_before=0, line=1.3, cell=None):
        p = (cell or doc).add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.line_spacing = line
        return p

    def run(p, text, size=9, color=BODY, bold=False):
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.bold = bold
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        return r

    def section(title, en):
        p = para(space_after=4, space_before=9)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "D9D9D9")
        pBdr.append(bottom)
        pPr.append(pBdr)
        run(p, title, size=11.5, color=INK, bold=True)
        run(p, "   " + en, size=7, color=RGBColor(0xA8, 0xA8, 0xA8))

    def right_row(left, right, left_size=10.5, left_bold=True):
        p = para(space_after=1)
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)
        run(p, left, size=left_size, color=INK, bold=left_bold)
        run(p, "\t" + right, size=8.5, color=GRAY)
        return p

    def bullet(label, text):
        p = para(space_after=1)
        p.paragraph_format.left_indent = Cm(0.42)
        p.paragraph_format.first_line_indent = Cm(-0.42)
        run(p, "·  ", size=9, color=GRAY)
        run(p, label, size=9, color=INK, bold=True)
        run(p, "：" + text, size=9)
        return p

    # 页首：左信息 + 右证件照（无边框表格承载）
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    cell_l, cell_r = table.cell(0, 0), table.cell(0, 1)
    cell_l.width, cell_r.width = Cm(14.6), Cm(3.4)
    table.columns[0].width, table.columns[1].width = Cm(14.6), Cm(3.4)
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        borders.append(el)
    tblPr.append(borders)
    cellMar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), "0")
        el.set(qn("w:type"), "dxa")
        cellMar.append(el)
    tblPr.append(cellMar)
    cell_l.paragraphs[0]._element.getparent().remove(cell_l.paragraphs[0]._element)

    p = para(cell=cell_l)
    run(p, DATA["name"], size=24, color=INK, bold=True)
    p = para(space_after=1, cell=cell_l)
    run(p, "求职意向 · " + DATA["intent"], size=9.5, color=INK)
    meta = "   ·   ".join(DATA["meta"])
    p = para(space_after=2, cell=cell_l)
    run(p, meta, size=8.5, color=GRAY)
    p = para(space_after=3, cell=cell_l)
    run(p, "   ·   ".join(DATA["contact"]), size=8.5, color=RGBColor(0x44, 0x44, 0x44))

    photo = ROOT / "photo.png"
    if DATA.get("show_photo", True) and photo.exists():
        rp = para(cell=cell_r)
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp.add_run().add_picture(str(photo), width=Cm(2.5))

    section("自我概述", "SUMMARY")
    run(para(space_after=0), DATA["summary"], size=9)

    section("核心能力", "COMPETENCIES")
    for c in DATA["competencies"]:
        p = para(space_after=1)
        p.paragraph_format.left_indent = Cm(0.42)
        p.paragraph_format.first_line_indent = Cm(-0.42)
        run(p, c["label"], size=9, color=INK, bold=True)
        run(p, "　　" + c["text"], size=9)

    section("工作经历", "EXPERIENCE")
    for e in DATA["experience"]:
        right_row(e["company"], e["date"])
        p = para(space_after=2)
        run(p, f"{e['role']}  /  {e['industry']}", size=8.8, color=RGBColor(0x55, 0x55, 0x55))
        for b in e["bullets"]:
            bullet(b["label"], b["text"])
        para(space_after=0, line=0.4).add_run("")

    section("教育经历", "EDUCATION")
    for e in DATA["education"]:
        right_row(e["school"], e["date"])
        p = para(space_after=1)
        run(p, f"{e['major']}  /  {e['degree']}", size=8.8, color=RGBColor(0x55, 0x55, 0x55))
        run(para(space_after=0), e["detail"], size=9)
    p = para(space_after=0, space_before=5)
    p.paragraph_format.left_indent = Cm(0.42)
    p.paragraph_format.first_line_indent = Cm(-0.42)
    run(p, "证书资质", size=9, color=INK, bold=True)
    run(p, "　　" + "  |  ".join(DATA["certificates"]), size=9)

    doc.save(DOCX_OUT)


def main():
    ap = argparse.ArgumentParser(description="简历生成：PDF / PNG / Word")
    ap.add_argument("step", nargs="?", default="all",
                    choices=["all", "pdf", "png", "docx"], help="生成步骤")
    ap.add_argument("--root", default=".", help="简历内容目录（含 resume.yml / photo.png）")
    ap.add_argument("--out", help="产物目录，默认与 --root 相同")
    a = ap.parse_args()
    configure(a.root, a.out)

    for layout, palette, font in MATCHES:
        html = render_html(layout, palette, font)
        print("html", html)
        if a.step in ("all", "pdf"):
            make_pdf(html, pdf_path(layout, palette, font))
            print("pdf", pdf_path(layout, palette, font))
        if a.step in ("all", "png"):
            make_png(html, layout, palette, font)
    if a.step in ("all", "docx"):
        make_docx()
        print("docx", DOCX_OUT)


if __name__ == "__main__":
    main()
